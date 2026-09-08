from __future__ import annotations

import logging
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional, Set, Tuple

from flask import jsonify, render_template, session
from weasyprint import HTML
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from config import BASE_DIR
from db import (
    fetch_departments,
    fetch_group_report_list,
    fetch_plant_list_by_group,
    fetch_purchasing_units,
    fetch_rows,
)
from helpers import (
    format_currency,
    format_thai_short_date,
    get_arg,
    insert_thai_break_hints,
    safe_fetch,
    safe_text,
    thai_month_name,
    to_buddhist_year,
    to_int_or_none,
    parse_month,
)

logger = logging.getLogger(__name__)

# ===================== K2 NAME/CODE RESOLUTION =====================
def resolve_dept_id_by_name(name: str) -> Optional[str]:
    """แปลงชื่อหน่วยงานจาก K2 (CUR_INDEPT_LNAME_TH) -> Dept_id ด้วย exact match"""
    return resolve_dept_by_name(name)[0]


def resolve_dept_by_name(name: str) -> Tuple[Optional[str], Optional[str]]:
    """
    คืน (Dept_id, Dept_name) จาก DB ถ้า match ชื่อที่ K2 ส่งมาแบบ exact เป๊ะ, (None, None) ถ้าไม่เจอ
    สำคัญ: ใช้ Dept_name ที่คืนมา (ไม่ใช่ข้อความดิบจาก K2) ไปโชว์หัวรายงาน เพื่อการันตีว่าชื่อที่โชว์
    ตรงกับข้อมูลที่ถูกกรองจริงเสมอ — ถ้า resolve ไม่เจอ ห้ามเอาข้อความดิบไปโชว์เด็ดขาด (จะโชว์ชื่อที่ไม่มีจริง)
    """
    if not name:
        return None, None
    name_norm = name.strip()
    for d in safe_fetch(lambda: fetch_departments(True), []):
        if d["name"].strip() == name_norm:
            return d["id"], d["name"]
    logger.warning("resolve_dept_by_name: ไม่พบหน่วยงาน '%s'", name_norm)
    return None, None


def resolve_plant_by_name(name: str) -> Optional[str]:
    """แปลงชื่อ Plant จาก K2 (PlantName) -> รหัส Plant"""
    if not name:
        return None
    name_norm = name.strip()
    for u in safe_fetch(fetch_purchasing_units, []):
        if u["name"].strip() == name_norm:
            return u["id"]
    logger.warning("resolve_plant_by_name: ไม่พบหน่วยจัดซื้อ '%s'", name_norm)
    return None


def resolve_group_report_by_label(text: str) -> Optional[str]:
    """
    แปลงข้อความ Group_Report_Desc (เช่น 'สำนักงานอธิการบดี', 'คณะ/สถาบัน/สำนัก') -> Group_Report code
    ถ้า K2 ส่งข้อความมาแทน code ตรง ๆ, หรือส่งข้อความที่ไม่ตรงกับ Group_Report_ms เลย (เช่น
    'มหาวิทยาลัยศรีนครินทรวิโรฒ' ซึ่งหมายถึง "ทั้งหมด" ไม่ใช่กลุ่มใดกลุ่มหนึ่ง) จะคืน None -> ไม่กรอง (=ทั้งมหาวิทยาลัย)
    """
    if not text:
        return None
    text_norm = text.strip()
    for g in safe_fetch(fetch_group_report_list, []):
        if g.get("label", "").strip() == text_norm:
            return g.get("value")
    logger.info("resolve_group_report_by_label: ไม่พบกลุ่ม '%s' ใน Group_Report_ms ถือว่าไม่กรอง (ทั้งมหาวิทยาลัย)", text_norm)
    return None


def extract_plant_code(row: Dict[str, Any]) -> Optional[str]:
    """ลองหา Plant จากแถว SP_4001 — เช็คด้วย /debug/sp4001-columns ว่าคอลัมน์จริงชื่ออะไร"""
    for key in ("Plant", "PlantCode", "plant", "Plant_id"):
        if row.get(key):
            return safe_text(row.get(key))
    return None


def filter_rows_by_plant(rows: List[Dict[str, Any]], plant_filter: Optional[str], allowed_plants: Optional[Set[str]]):
    """กรอง rows ตาม Plant เดียวหรือกลุ่ม Plant ฝั่ง Python (SP ไม่รองรับกรองแบบกลุ่ม)"""
    if not plant_filter and not allowed_plants:
        return rows
    filtered = []
    for r in rows:
        plant_code = extract_plant_code(r)
        if plant_code is None:
            continue  # หา Plant ไม่ได้ -> ตัดออก ปลอดภัยกว่าเดา
        if (plant_filter and plant_code == plant_filter) or (allowed_plants and plant_code in allowed_plants):
            filtered.append(r)
    return filtered


def is_university_wide_group(group_report: Optional[str]) -> bool:
    """0 = มหาวิทยาลัย (ทั้งหมด), 1 = สนอ., 2 = คณะ — เช็ค '0' หลัก, fallback เช็ค label"""
    if group_report is None:
        return False
    if str(group_report) == "0":
        return True
    found = next((g for g in safe_fetch(fetch_group_report_list, []) if g.get("value") == group_report), None)
    return bool(found and "มหาวิทยาลัย" in (found.get("label") or ""))


def resolve_unit_display_name(dept_id, sp_plant_param, group_report, rows) -> str:
    """หัวรายงาน: ลำดับ Plant เจาะจง > Dept เจาะจง > กลุ่มที่เลือก > ว่าง (=ทั้งมหาวิทยาลัย)"""
    if sp_plant_param:
        found = next((u for u in safe_fetch(fetch_purchasing_units, []) if str(u.get("id")) == str(sp_plant_param)), None)
        if found and found.get("name"):
            return found["name"]

    if dept_id and rows:
        name = safe_text(rows[0].get("Dept_name"))
        if name != "-":
            return name

    if group_report and not is_university_wide_group(group_report):
        found = next((g for g in safe_fetch(fetch_group_report_list, []) if g.get("value") == group_report), None)
        if found and found.get("label"):
            return found["label"]

    return ""


# ===================== FILTER RESOLUTION / SCOPING =====================
def resolve_query_filters():
    """parse + resolve query params ที่ใช้ร่วมกันในทั้ง pdf และ xlsx route"""
    fiscal_year = to_buddhist_year(to_int_or_none(get_arg("FiscalYear", "fiscal_year", "Year")))
    # ใช้เดือนที่ผู้ใช้เลือกตรงๆ ทั้งกรองข้อมูลและตั้งชื่อเดือนในหัวรายงาน (ไม่ถอยหลัง 1 เดือนอีกต่อไป —
    # เดือน 8 ที่เลือก ต้องได้ทั้งข้อมูลของเดือน 8 และหัวรายงานเขียนว่า "เดือนสิงหาคม" ตรงกัน)
    month = parse_month(get_arg("Month", "month"))
    dept_id = get_arg("Search_Dept_ID", "DeptID", "dept_id", default="").strip() or None
    dept_name = get_arg("CUR_INDEPT_LNAME_TH", "DeptName", default="").strip() or None
    plant_id_direct = get_arg("PlantName_id", "PlantID", "Search_Plant", "Plant", default="").strip() or None
    plant_name = get_arg("PlantName", default="").strip() or None
    group_report_raw = get_arg("Group_Report", "ReportGroup", default="").strip() or None

    if dept_id and "," in dept_id:
        logger.warning("Search_Dept_ID เป็น comma list ('%s') ไม่รองรับ ตัดทิ้ง", dept_id)
        dept_id = None
    resolved_dept_name = None  # ชื่อหน่วยงานจริงจาก DB ที่ยืนยันแล้วว่าตรงกับ dept_id — ใช้โชว์หัวรายงานได้อย่างปลอดภัยเท่านั้น
    if not dept_id and dept_name:
        dept_id, resolved_dept_name = resolve_dept_by_name(dept_name)

    plant_filter = plant_id_direct or (resolve_plant_by_name(plant_name) if plant_name else None)
    try:
        sp_plant_param = int(plant_filter) if plant_filter else None
    except (TypeError, ValueError):
        sp_plant_param = None

    # Group_Report: ยอมรับได้ทั้ง code ตรงๆ (เช่น "1") หรือข้อความ label (เช่น "สำนักงานอธิการบดี")
    # ถ้าเป็น code ที่มีอยู่จริงใน Group_Report_ms ใช้เลย, ถ้าไม่ใช่ลองแปลงจาก label, ถ้าแปลงไม่ได้ = ไม่กรอง (ทั้งมหาวิทยาลัย)
    group_report = None
    if group_report_raw:
        known_codes = {g.get("value") for g in safe_fetch(fetch_group_report_list, [])}
        group_report = group_report_raw if group_report_raw in known_codes else resolve_group_report_by_label(group_report_raw)

    return fiscal_year, month, dept_id, sp_plant_param, group_report, resolved_dept_name


def fetch_filtered_rows(fiscal_year, month, dept_id, sp_plant_param, group_report):
    """ดึงจาก SP_4001 แล้วกรองตามกลุ่ม (Group_Report) เพิ่มถ้าจำเป็น — SP กรองได้แค่ Plant เดียว"""
    rows = fetch_rows(fiscal_year, month, dept_id, sp_plant_param)
    if group_report and not sp_plant_param:
        allowed = safe_fetch(lambda: set(fetch_plant_list_by_group(group_report)), None)
        rows = filter_rows_by_plant(rows, None, allowed)
    return rows


def scoped_fetch_filtered_rows(fiscal_year, month, dept_id, sp_plant_param, group_report):
    """
    ตัวห่อ fetch_filtered_rows ที่บังคับสิทธิ์ตาม session (สำคัญ: กัน user ที่ไม่ใช่ admin
    แก้ query string เอง เช่น Search_Dept_ID/Search_Plant เป็นของหน่วยงานอื่น แล้วดึงข้อมูลนอกสิทธิ์ตนเอง)
    คืนค่า (rows, error_response_or_None) — ถ้ามี error_response ให้ return ค่านั้นตรงๆ จาก route
    """
    if session.get("role") == "admin":
        return fetch_filtered_rows(fiscal_year, month, dept_id, sp_plant_param, group_report), None

    allowed_depts = set(str(d) for d in (session.get("dept_ids") or []))
    allowed_plants = set(str(p) for p in (session.get("plants") or []))

    if dept_id is not None and str(dept_id) not in allowed_depts:
        return None, (jsonify({"error": "ไม่มีสิทธิ์เข้าถึงข้อมูลของหน่วยงานนี้"}), 403)
    if sp_plant_param is not None and str(sp_plant_param) not in allowed_plants:
        return None, (jsonify({"error": "ไม่มีสิทธิ์เข้าถึงข้อมูลของหน่วยจัดซื้อนี้"}), 403)

    rows = fetch_filtered_rows(fiscal_year, month, dept_id, sp_plant_param, group_report)
    if dept_id is None and sp_plant_param is None:
        # ไม่ได้เจาะจงหน่วยงาน/หน่วยจัดซื้อมา -> จำกัดผลลัพธ์ให้เหลือเฉพาะที่ตนเองมีสิทธิ์เท่านั้น
        # (ไม่งั้น user ทั่วไปจะเห็นข้อมูลทั้งมหาวิทยาลัยได้ถ้าไม่ระบุ filter)
        rows = filter_rows_by_plant(rows, None, allowed_plants)
    return rows, None


def current_report_date() -> Tuple[str, str, str]:
    """
    วันที่ออกรายงาน (แถว 'วันที่ ... เดือน ... พ.ศ. ...' บนหัวเอกสาร) ใช้วันที่ปัจจุบันของเซิร์ฟเวอร์
    เสมอ ไม่รับ override จาก query string/K2 อีกต่อไป — เพราะวันที่นี้หมายถึง "วันที่พิมพ์รายงาน"
    ไม่ใช่ช่วงเวลาของข้อมูล (ช่วงเวลาของข้อมูลควบคุมด้วย FiscalYear/Month แยกต่างหาก)
    """
    today = datetime.now()
    return str(today.day), thai_month_name(today.month), str(today.year + 543)


# ===================== DATA MAPPING / RENDERING =====================
def map_row(row: Dict[str, Any]) -> Dict[str, str]:
    contract_date_str = format_thai_short_date(row.get("RefDocNoDate"))

    return {
        "seq": safe_text(row.get("SeqNo")),
        "work": insert_thai_break_hints(safe_text(row.get("WorkDetail"))),
        "budget": format_currency(row.get("Budget")),
        "median": format_currency(row.get("MedianPrice")),
        "method": insert_thai_break_hints(safe_text(row.get("Method_name"))),
        "bidder_list": insert_thai_break_hints(safe_text(row.get("BidderList"))),
        "winner": insert_thai_break_hints(safe_text(row.get("Winner"))),
        "reason": insert_thai_break_hints(safe_text(row.get("Reason_name"))),
        "refdoc": safe_text(row.get("RefDocNo")),
        "contract_date": contract_date_str,
        "dept_name": safe_text(row.get("Dept_name")) if row.get("Dept_name") else safe_text(row.get("Dept_id")),
    }


def render_pdf_bytes(rows, month_name, report_year, unit_name, report_day, report_month, report_be_year) -> bytes:
    """PDF ด้วย WeasyPrint (เสถียรกว่า Playwright/Chromium กับหน้าข้อมูลเยอะ)"""
    html_content = render_template(
        "pdf_report.html",
        rows=[map_row(r) for r in rows],
        month_name=month_name, report_year=report_year, unit_name=unit_name,
        report_day=report_day, report_month=report_month, report_be_year=report_be_year,
    )
    return HTML(string=html_content, base_url=str(BASE_DIR)).write_pdf()


# หัวคอลัมน์ตรงกับ pdf_report.html เป๊ะๆ — คอลัมน์สุดท้าย (เลขที่/วันที่สัญญา) merge เป็นหัวเดียว colspan=2 เหมือน PDF
# หมายเหตุ: PDF ไม่มีคอลัมน์ "หน่วยงาน" (หัวรายงานระบุหน่วยงานไว้ที่ subtitle ด้านบนแทน) จึงตัดออกจาก xlsx ด้วยเพื่อให้ตรงกัน
REPORT_COLUMN_HEADERS = [
    "ลำดับ",
    "งานที่จัดซื้อหรือจัดจ้าง",
    "วงเงินที่จะซื้อหรือจ้าง (บาท)",
    "ราคากลาง (บาท)",
    "วิธีซื้อหรือจ้าง",
    "รายชื่อผู้เสนอราคาและราคาที่เสนอ",
    "ผู้ได้รับการคัดเลือกและราคาที่ตกลงซื้อหรือจ้าง",
    "เหตุผลที่คัดเลือกโดยสรุป",
]
CONTRACT_HEADER = "เลขที่และวันที่ของสัญญาหรือข้อตกลงในการซื้อหรือจ้าง"  # merge คอลัมน์ 9-10 เหมือน <th colspan="2"> ใน PDF


def render_xlsx_bytes(
    rows: List[Dict[str, Any]],
    month_name: str,
    fiscal_year: Optional[int],
    unit_name: str,
    report_day: str,
    report_month: str,
    report_be_year: str,
) -> bytes:
    """
    สร้างไฟล์ .xlsx ด้วย openpyxl — ข้อความหัวเรื่อง/หน่วยงาน/วันที่ และหัวคอลัมน์ ใช้ประโยคเดียวกับ pdf_report.html
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "รายงาน"

    col_count = len(REPORT_COLUMN_HEADERS) + 2  # 8 หัวเดี่ยว + 2 คอลัมน์ที่ merge กันเป็นหัวเดียว = 10 คอลัมน์ข้อมูล

    title_font = Font(name="TH Sarabun New", size=16, bold=True)
    subtitle_font = Font(name="TH Sarabun New", size=14)
    header_font = Font(name="TH Sarabun New", size=14, bold=True)
    body_font = Font(name="TH Sarabun New", size=14)
    header_fill = PatternFill(start_color="EBEBEB", end_color="EBEBEB", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    right_top = Alignment(horizontal="right", vertical="top")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row_idx = 1

    # เลขแบบฟอร์มมุมขวาบน — เทียบเท่า <div class="form-code">สขร.1</div> ใน pdf_report.html
    # ไม่ merge เซลล์นี้ (ต่างจากหัวเรื่อง/หน่วยงาน/วันที่) เพราะต้องการแค่ชิดขวาสุดคอลัมน์สุดท้าย
    code_cell = ws.cell(row=row_idx, column=col_count, value="สขร.1")
    code_cell.font = subtitle_font
    code_cell.alignment = Alignment(horizontal="right", vertical="center")
    row_idx += 1

    # หัวเรื่อง — ตรงกับ <div class="title"> ใน pdf_report.html
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=col_count)
    title_cell = ws.cell(row=row_idx, column=1, value=f"สรุปผลการดำเนินการจัดซื้อจัดจ้างในรอบเดือน{month_name} {fiscal_year or ''}")
    title_cell.font = title_font
    title_cell.alignment = center
    row_idx += 1

    # หน่วยงาน — ตรงกับ <div class="subtitle"> (unit_name + มหาวิทยาลัยศรีนครินทรวิโรฒ)
    subtitle = f"{unit_name} มหาวิทยาลัยศรีนครินทรวิโรฒ" if unit_name else "มหาวิทยาลัยศรีนครินทรวิโรฒ"
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=col_count)
    subtitle_cell = ws.cell(row=row_idx, column=1, value=subtitle)
    subtitle_cell.font = subtitle_font
    subtitle_cell.alignment = center
    row_idx += 1

    # วันที่ออกรายงาน — ตรงกับ <div class="date-line">
    if report_day and report_month and report_be_year:
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=col_count)
        date_cell = ws.cell(row=row_idx, column=1, value=f"วันที่ {report_day} เดือน{report_month} พ.ศ. {report_be_year}")
        date_cell.font = subtitle_font
        date_cell.alignment = center
        row_idx += 1

    row_idx += 1  # แถวว่างคั่นก่อนตาราง

    # หัวตาราง — 8 คอลัมน์เดี่ยว + 1 หัว merge คู่ (เลขที่และวันที่ของสัญญา) เหมือน PDF
    header_row = row_idx
    for col, text in enumerate(REPORT_COLUMN_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    last_col = len(REPORT_COLUMN_HEADERS)
    ws.merge_cells(start_row=header_row, start_column=last_col + 1, end_row=header_row, end_column=last_col + 2)
    for col in (last_col + 1, last_col + 2):
        cell = ws.cell(row=header_row, column=col, value=CONTRACT_HEADER if col == last_col + 1 else None)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    row_idx += 1

    # แถวข้อมูล — ลำดับคอลัมน์: ลำดับ, งาน, วงเงิน, ราคากลาง, วิธี, ผู้เสนอราคา, ผู้ชนะ, เหตุผล, เลขที่สัญญา, วันที่สัญญา
    numeric_cols = {3, 4}       # วงเงิน, ราคากลาง -> ชิดขวา
    center_cols = {1, 5, 9, 10}  # ลำดับ, วิธีซื้อหรือจ้าง, เลขที่สัญญา, วันที่สัญญา -> กึ่งกลาง
    data_start_row = row_idx
    for row in rows:
        # หมายเหตุ: คู่กับ RefDocNo คือ RefDocNoDate (ไม่ใช่ ContractDate) — เหมือนที่แก้ใน map_row สำหรับ PDF
        contract_date_str = format_thai_short_date(row.get("RefDocNoDate"))
        values = [
            safe_text(row.get("SeqNo")), safe_text(row.get("WorkDetail")),
            safe_text(row.get("Budget")), safe_text(row.get("MedianPrice")), safe_text(row.get("Method_name")),
            safe_text(row.get("BidderList")), safe_text(row.get("Winner")), safe_text(row.get("Reason_name")),
            safe_text(row.get("RefDocNo")), contract_date_str,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.font = body_font
            cell.border = border
            if col in numeric_cols:
                cell.alignment = right_top
            elif col in center_cols:
                cell.alignment = center
            else:
                cell.alignment = left_top
        row_idx += 1

    last_data_row = row_idx - 1

    # ความกว้างคอลัมน์โดยประมาณ (หน่วย: จำนวนตัวอักษร) — 10 คอลัมน์ตรงกับข้อมูล
    column_widths = [6, 30, 14, 14, 16, 26, 26, 20, 16, 12]
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = ws.cell(row=data_start_row, column=1).coordinate

    # ===== ตั้งค่าหน้าพิมพ์ — ให้เปิดใน Excel แล้วกด Print/Save as PDF ได้หน้าสวย ไม่ตัดคอลัมน์กลางตาราง =====
    # A4 แนวนอน (เหมือนหน้า PDF ต้นฉบับ), บีบให้พอดี 1 หน้าตามความกว้าง แต่ปล่อยความสูงอัตโนมัติ
    # (fitToHeight=0) เพื่อให้ยาวได้หลายหน้าตามจำนวนแถวข้อมูลโดยไม่บีบตัวอักษรจนอ่านไม่ออก
    last_col_letter = get_column_letter(col_count)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)
    ws.print_options.horizontalCentered = True
    # พื้นที่พิมพ์ครอบทั้งตาราง (รวมแถว สขร.1/หัวเรื่อง/หน่วยงาน/วันที่ ด้านบนด้วย)
    ws.print_area = f"A1:{last_col_letter}{last_data_row}"
    # แถวหัวตาราง (ลำดับ/งานที่จัดซื้อ/ฯลฯ) พิมพ์ซ้ำทุกหน้า ให้รู้ว่าแต่ละคอลัมน์คือคอลัมน์อะไรแม้ขึ้นหน้าใหม่
    ws.print_title_rows = f"{header_row}:{header_row}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def report_filename(month_name: str, fiscal_year, ext: str) -> str:
    return f"สรุปผลการดำเนินการจัดซื้อจัดจ้างในรอบเดือน{month_name} {fiscal_year or ''}.{ext}"